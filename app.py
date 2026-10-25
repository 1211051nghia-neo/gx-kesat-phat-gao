import os
import io
import time
import requests
import pandas as pd
from datetime import datetime
from fastapi import FastAPI, HTTPException
from fastapi.responses import HTMLResponse, FileResponse, StreamingResponse
import re
from pydantic import BaseModel
from pathlib import Path

# Cố định đường dẫn theo vị trí của file
BASE_DIR = Path(__file__).resolve().parent
BG_IMAGE = BASE_DIR / "nha-tho-ke-sat-xl.JPG"

# ==========================================
# CẤU HÌNH API GOOGLE APPS SCRIPT CỦA BẠN
# ==========================================
GAS_URL = "https://script.google.com/macros/s/AKfycbwXVzN0-9F1AAV88zgp-dI3SK7-t4WwLtoaFvMwn7FS3SFV_bz7hpEMpviJGoKr7KPF/exec"

app = FastAPI(title="Gx. Kẻ Sặt phát gạo (Cloud Edition)")

# ==========================================
# CÁC HÀM GIAO TIẾP VỚI GOOGLE SHEETS
# ==========================================

# Bộ nhớ tạm (Cache) để giảm số lần gọi Google Sheets API
DATA_CACHE = {"data": ([], []), "timestamp": 0}
CACHE_TTL = 10  # Lưu dữ liệu trong 10 giây, sau 10s mới gọi lại Google Sheets

def fetch_all_data(force_refresh=False):
    """Hàm lấy toàn bộ dữ liệu từ Google Sheets về Python (Có Cache)"""
    now = time.time()
    if not force_refresh and (now - DATA_CACHE["timestamp"] < CACHE_TTL):
        return DATA_CACHE["data"]

    try:
        response = requests.get(f"{GAS_URL}?action=get_data", timeout=15)
        data = response.json()

        raw_ben = data.get("beneficiaries", [])
        beneficiaries = [r for r in raw_ben[1:] if any(r)] if len(raw_ben) > 1 else []
        
        raw_chk = data.get("checkins", [])
        checkins = [r for r in raw_chk[1:] if any(r)] if len(raw_chk) > 1 else []
        
        DATA_CACHE["data"] = (beneficiaries, checkins)
        DATA_CACHE["timestamp"] = now
        return beneficiaries, checkins
    
    except Exception as e:
        print(f"Lỗi kết nối Google Sheets: {e}")
        return DATA_CACHE["data"] # Trả về cache cũ nếu mất mạng






    




    

def send_post_action(action, payload):
    """Hàm gửi lệnh Thêm/Sửa/Xóa lên Google Sheets"""
    payload["action"] = action
    try:
        requests.post(GAS_URL, json=payload, timeout=10)
    except Exception as e:
        print(f"Lỗi gửi dữ liệu lên Google Sheets: {e}")



# ==========================================
# SCHEMAS
# ==========================================
class CheckinRequest(BaseModel):
    stt: int
    collaborator_name: str = "Cộng tác viên"

class BeneficiarySchema(BaseModel):
    full_name: str
    address: str = ""
    job: str = ""
    notes: str = ""

# ==========================================
# CÁC API ENDPOINTS
# ==========================================

@app.get("/static/background.jpg")
def get_background():
    if Path(BG_IMAGE).exists():
        return FileResponse(BG_IMAGE)
    return HTMLResponse(status_code=404)


@app.get("/api/beneficiaries")
def get_all_beneficiaries():
    current_month = datetime.now().strftime("%Y-%m")
    beneficiaries, checkins = fetch_all_data()
    
    # Tab Checkins: chk[0]=ID, chk[1]=STT, chk[2]=Thang, chk[3]=ThoiGian, chk[4]=CTV
    # Lọc danh sách checkin trong tháng hiện tại
    current_checkins = {}
    for chk in checkins:
        if len(chk) > 2:
            # Lấy 7 ký tự đầu (YYYY-MM) tránh lỗi định dạng Date của Excel/GSheets
            chk_month = str(chk[2]).strip()[:7]
            if chk_month == current_month:
                current_checkins[str(chk[1]).strip()] = chk
    
    result = []
    for b in beneficiaries:
        if not b or len(b) == 0 or not str(b[0]).strip():
            continue
        try:
            stt_val = int(b[0])
        except:
            continue
            
        stt_str = str(stt_val)
        chk = current_checkins.get(stt_str)
        
        result.append({
            "stt": stt_val,
            "full_name": str(b[1]) if len(b) > 1 else "",
            "address": str(b[2]) if len(b) > 2 else "",
            "job": str(b[3]) if len(b) > 3 else "",
            "notes": str(b[4]) if len(b) > 4 else "",
            "avatar_url": f"https://ui-avatars.com/api/?name={str(b[1] if len(b) > 1 else 'User').replace(' ', '+')}&background=random",
            "received_at": str(chk[3]) if chk and len(chk) > 3 else None,
            "collaborator_name": str(chk[4]) if chk and len(chk) > 4 else None
        })
    return result




@app.get("/api/beneficiary/{code}")
def get_beneficiary(code: str):
    current_month = datetime.now().strftime("%Y-%m")
    clean_code = code.strip().lower()
    
    beneficiaries, checkins = fetch_all_data()
    person = None

    # Logic tìm kiếm tương tự SQL nhưng dùng vòng lặp Python
    # TRƯỜNG HỢP 1: Mã QR CCCD (có chứa dấu |)
    if "|" in clean_code:
        parts = [p.strip() for p in clean_code.split("|")]
        if len(parts) >= 3 and parts[2]:
            scanned_name = parts[2].lower()
            for b in beneficiaries:
                if scanned_name == str(b[1]).strip().lower():
                    person = b
                    break

    # TRƯỜNG HỢP 2 & 3: Có chứa số (Mã GAO hoặc gõ STT)
    elif re.search(r'\d+', clean_code):
        match = re.search(r'\d+', clean_code)
        if match:
            stt_val = str(int(match.group()))
            for b in beneficiaries:
                if str(b[0]) == stt_val:
                    person = b
                    break

    # TRƯỜNG HỢP 4: Gõ tay một phần tên
    if not person:
        for b in beneficiaries:
            if clean_code in str(b[1]).strip().lower():
                person = b
                break

    if not person:
        raise HTTPException(status_code=404, detail=f"Không tìm thấy dữ liệu cho mã: '{code}'")

    stt_val = int(person[0])
    stt_str = str(stt_val)
    
    # Kiểm tra đã nhận tháng này chưa (Sửa so sánh 7 ký tự đầu YYYY-MM)
    checkin = None
    for chk in checkins:
        if len(chk) > 2:
            chk_stt = str(chk[1]).strip()
            chk_month = str(chk[2]).strip()[:7]
            if chk_stt == stt_str and chk_month == current_month:
                checkin = chk
                break


    full_name = str(person[1]) if len(person) > 1 else ""
    return {
        "stt": stt_val,
        "full_name": full_name,
        "address": str(person[2]) if len(person) > 2 else "",
        "job": str(person[3]) if len(person) > 3 else "",
        "notes": str(person[4]) if len(person) > 4 else "",
        "avatar_url": f"https://ui-avatars.com/api/?name={full_name.replace(' ', '+')}&background=random",
        "already_received": checkin is not None,
        "received_at": str(checkin[3]) if checkin and len(checkin) > 3 else None,
        "collaborator_name": str(checkin[4]) if checkin and len(checkin) > 4 else None,
    }




@app.post("/api/beneficiary")
def add_beneficiary(data: BeneficiarySchema):
    beneficiaries, _ = fetch_all_data()
    # Tính STT mới
    max_stt = 0
    for b in beneficiaries:
        try:
            val = int(b[0])
            if val > max_stt: max_stt = val
        except: pass
    
    new_stt = max_stt + 1
    send_post_action("add_beneficiary", {
        "stt": new_stt, "full_name": data.full_name, 
        "address": data.address, "job": data.job, "notes": data.notes
    })
    return {"message": f"Thêm hộ mới thành công (STT: #{new_stt})!", "stt": new_stt}


@app.put("/api/beneficiary/{stt}")
def update_beneficiary(stt: int, data: BeneficiarySchema):
    send_post_action("update_beneficiary", {
        "stt": stt, "full_name": data.full_name, 
        "address": data.address, "job": data.job, "notes": data.notes
    })
    return {"message": "Cập nhật thông tin thành công!"}


@app.post("/api/checkin")
def do_checkin(req: CheckinRequest):
    current_month = datetime.now().strftime("%Y-%m")
    now_str = datetime.now().strftime("%H:%M:%S - %d/%m/%Y")
    
    # Gọi Google Sheet để ghi dữ liệu (ID random sinh bằng Timestamp)
    random_id = int(time.time() * 1000)
    send_post_action("checkin", {
        "id": random_id,
        "stt": req.stt,
        "thang": current_month,
        "thoigian": now_str,
        "ctv": req.collaborator_name
    })

    # Xóa Cache để lượt quét tiếp theo tự làm mới từ Google Sheet
    DATA_CACHE["timestamp"] = 0

    return {"message": "Xác nhận phát gạo thành công!", "received_at": now_str}


@app.delete("/api/checkin/{stt}")
def delete_checkin(stt: int):
    current_month = datetime.now().strftime("%Y-%m")
    send_post_action("delete_checkin", {"stt": stt, "thang": current_month})
    return {"message": "Đã thu hồi lịch sử nhận gạo thành công!"}


@app.get("/api/checkins/history")
def get_history():
    current_month = datetime.now().strftime("%Y-%m")
    beneficiaries, checkins = fetch_all_data()
    
    ben_map = {}
    for b in beneficiaries:
        if b and len(b) > 0 and str(b[0]).strip():
            try:
                stt_key = str(int(float(b[0])))
                ben_map[stt_key] = b
            except:
                pass
            
    result = []
    for chk in reversed(checkins):
        if len(chk) > 2:
            chk_month = str(chk[2]).strip()[:7]
            if chk_month == current_month:
                try:
                    stt_str = str(int(float(chk[1])))
                except:
                    stt_str = str(chk[1]).strip()
                    
                person = ben_map.get(stt_str, ["", "Không rõ", "", "", ""])
                
                months_received = [str(c[2])[:7] for c in checkins if len(c) > 2 and str(c[1]).strip() == stt_str]
                months_str = ", ".join(list(set(months_received)))
                
                try:
                    chk_id = int(float(chk[0]))
                except:
                    chk_id = int(time.time() * 1000)

                result.append({
                    "id": chk_id,
                    "created_at": str(chk[3]) if len(chk) > 3 else "",
                    "collaborator_name": str(chk[4]) if len(chk) > 4 else "CTV",
                    "stt": int(float(person[0])) if person[0] != "" else 0,
                    "full_name": str(person[1]) if len(person) > 1 else "Không rõ",
                    "address": str(person[2]) if len(person) > 2 else "",
                    "job": str(person[3]) if len(person) > 3 else "",
                    "notes": str(person[4]) if len(person) > 4 else "",
                    "months_received": months_str
                })
    return result





# ==================== 2 API XUẤT EXCEL BẰNG PANDAS ====================
@app.get("/api/export/beneficiaries")
def export_beneficiaries_excel():
    current_month = datetime.now().strftime("%Y-%m")
    beneficiaries, checkins = fetch_all_data()
    
    current_checkins = {str(chk[1]): chk for chk in checkins if str(chk[2])[:7] == current_month}
    
    rows = []
    for b in beneficiaries:
        stt = str(b[0])
        chk = current_checkins.get(stt)
        rows.append({
            "STT": b[0],
            "Họ và Tên": str(b[1]),
            "Địa Chỉ": str(b[2]),
            "Nghề Nghiệp": str(b[3]),
            "Ghi Chú": str(b[4]),
            "Trạng Thái Tháng Này": "Đã nhận" if chk else "Chưa nhận",
            "Thời Gian Nhận": chk[3] if chk else "",
            "CTV Thực Hiện": chk[4] if chk else ""
        })
    
    df = pd.DataFrame(rows)
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='DanhSachHoDan', index=False)
        
        # --- LÀM ĐẸP EXCEL (STYLING) ---
        workbook = writer.book
        worksheet = writer.sheets['DanhSachHoDan']
        
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        
        header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
        align_center = Alignment(horizontal='center', vertical='center')
        align_left = Alignment(horizontal='left', vertical='center')
        thin_border = Border(
            left=Side(style='thin', color='CCCCCC'),
            right=Side(style='thin', color='CCCCCC'),
            top=Side(style='thin', color='CCCCCC'),
            bottom=Side(style='thin', color='CCCCCC')
        )
        
        # Định dạng dòng tiêu đề (Header)
        for col_num in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center
            cell.border = thin_border
            
        # Định dạng các dòng dữ liệu
        for row_num in range(2, len(df) + 2):
            for col_num in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row_num, column=col_num)
                # Kiểm tra cột 2 (Họ và Tên) để in đậm
                if col_num == 2:
                    cell.font = Font(name='Segoe UI', size=10, bold=True)
                else:
                    cell.font = Font(name='Segoe UI', size=10)
                cell.border = thin_border
                if col_num in [1, 6]:  # Cột STT và Trạng Thái căn giữa
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left
                    
        # Tự động điều chỉnh độ rộng cột cho vừa khít nội dung
        for col in worksheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
    output.seek(0)
    filename = f"Danh_Sach_Ho_Dan_{current_month}.xlsx"
    headers = {'Content-Disposition': f'attachment; filename="{filename}"'}
    return StreamingResponse(output, headers=headers, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.get("/api/export/history")
def export_history_excel():
    current_month = datetime.now().strftime("%Y-%m")
    beneficiaries, checkins = fetch_all_data()
    
    ben_map = {str(b[0]): b for b in beneficiaries}
    
    rows = []
    for chk in reversed(checkins):
        if len(chk) > 2 and str(chk[2])[:7] == current_month:
            stt = str(chk[1])
            person = ben_map.get(stt, ["", "Không rõ", "", "", ""])
            months_received = [str(c[2])[5:7] for c in checkins if len(c) > 2 and str(c[1]) == stt]
            
            rows.append({
                "STT": person[0],
                "Họ và Tên": str(person[1]),
                "Địa Chỉ": str(person[2]),
                "Nghề Nghiệp": str(person[3]),
                "Ghi Chú": str(person[4]),
                "Các Tháng Đã Nhận": ", ".join(list(set(months_received))),
                "Thời Gian Nhận Gạo": str(chk[3]),
                "CTV Thực Hiện": str(chk[4])
            })
    
    df = pd.DataFrame(rows)
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='LichSuPhatGao', index=False)
        
        # --- LÀM ĐẸP EXCEL (STYLING) ---
        worksheet = writer.sheets['LichSuPhatGao']
        
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        
        header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
        align_center = Alignment(horizontal='center', vertical='center')
        align_left = Alignment(horizontal='left', vertical='center')
        thin_border = Border(
            left=Side(style='thin', color='CCCCCC'),
            right=Side(style='thin', color='CCCCCC'),
            top=Side(style='thin', color='CCCCCC'),
            bottom=Side(style='thin', color='CCCCCC')
        )

        # định dạng tiêu đề
        for col_num in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center
            cell.border = thin_border

        # định dạng các dòng    
        for row_num in range(2, len(df) + 2):
            for col_num in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row_num, column=col_num)
                # Kiểm tra cột 2 (Họ và Tên) để in đậm
                if col_num == 2:
                    cell.font = Font(name='Segoe UI', size=10, bold=True)
                else:
                    cell.font = Font(name='Segoe UI', size=10)
                cell.border = thin_border
                if col_num in [1, 6]:  # STT và Tháng đã nhận căn giữa
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left

                    
        for col in worksheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
    output.seek(0)
    filename = f"Lich_Su_Phat_Gao_{current_month}.xlsx"
    headers = {'Content-Disposition': f'attachment; filename="{filename}"'}
    return StreamingResponse(output, headers=headers, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')












# ==========================================
# 3. FRONTEND UI
# ==========================================
HTML_TEMPLATE = """
<!DOCTYPE html>
<html lang="vi" class="dark">

<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Gx. Kẻ Sặt thực thi bác ái</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <script src="https://unpkg.com/html5-qrcode"></script>
    <script>
        tailwind.config = {
            darkMode: 'class',
            theme: { extend: {} }
        }
    </script>

    <style>
        /* Thiết lập màu nền mặc định toàn trang */
        body { 
            margin: 0; 
            padding: 0; 
            background-color: #f1f5f9; /* Nền sáng dịu mắt */
            transition: background-color 0.3s ease;
        }
        .dark body { 
            background-color: #0b0f19; /* Nền tối sang trọng, sâu lắng, không bị đen ghê */
        }

        /* Hiệu ứng kính (Glassmorphism) thanh lịch, không dùng ảnh nền */
        .glass { 
            backdrop-filter: blur(16px); -webkit-backdrop-filter: blur(16px); 
            transition: all 0.3s;
        }
        .dark .glass { 
            background: rgba(17, 24, 39, 0.85); /* Xám xanh đậm, mờ nhẹ, cực kỳ dịu mắt */
            border: 1px solid rgba(255, 255, 255, 0.08);
            box-shadow: 0 10px 25px -5px rgba(0, 0, 0, 0.4);
        }
        html:not(.dark) .glass { 
            background: rgba(255, 255, 255, 0.9); 
            border: 1px solid rgba(0, 0, 0, 0.06);
            box-shadow: 0 10px 25px -5px rgba(0, 0, 0, 0.05); 
        }

        .dark-text { color: #1e293b; } .dark .dark-text { color: #f1f5f9; }
        .sub-text { color: #64748b; } .dark .sub-text { color: #94a3b8; }
        
        /* Ô input và select nổi bật, dễ nhìn ở cả 2 chế độ */
        .input-glass { background: rgba(255, 255, 255, 0.95); color: #0f172a; border: 1px solid rgba(0,0,0,0.12); }
        .dark .input-glass { background: rgba(15, 23, 42, 0.85); color: #f8fafc; border: 1px solid rgba(255,255,255,0.1); }

        #reader { border: none !important; border-radius: 1rem; overflow: hidden; background: transparent; }
        #reader video { border-radius: 1rem; object-fit: cover; }
        #reader__dashboard_section_csr span, #reader__dashboard_section_swaplink { color: inherit !important; font-size: 14px; margin-bottom:10px; }
        #reader button { background: #10b981 !important; color: white !important; border: none !important; padding: 8px 16px !important; border-radius: 8px !important; font-weight: bold; cursor: pointer; transition: 0.2s; }
        #reader button:hover { background: #059669 !important; }
        
        /* Ẩn các text tiếng Anh dư thừa của html5-qrcode
        #reader__dashboard_section_csr span { display: none !important; }
        #reader a { color: #10b981 !important; text-decoration: none; font-weight: bold; }
        */
    </style>



</head>

<body class="text-gray-900 dark:text-gray-100 min-h-screen pb-20 font-sans antialiased relative">
    

    <header class="sticky top-0 z-40 glass border-b border-gray-200 dark:border-gray-800 px-4 py-3 shadow-lg">
        <div class="max-w-7xl mx-auto flex flex-wrap justify-between items-center gap-2">
            <!--
            <h1 class="text-xl font-black text-transparent bg-clip-text bg-gradient-to-r from-green-500 to-blue-600 dark:from-green-400 dark:to-blue-500">
                "Chính anh em hãy cho họ ăn (Mt 14,16)"
            </h1>
            -->

            <h1 class="text-xl font-black text-yellow-300 text-center">
                "CHÍNH ANH EM HÃY CHO HỌ ĂN"<br/>(Mt 14,16) 
            </h1>

            <div class="flex items-center space-x-4">
                <label class="relative inline-flex items-center cursor-pointer" title="Đổi giao diện">
                    <input type="checkbox" id="theme-toggle" class="sr-only peer" onchange="toggleTheme()">
                    <div class="w-11 h-6 bg-gray-300 peer-focus:outline-none rounded-full peer dark:bg-gray-600 peer-checked:after:translate-x-full peer-checked:after:border-white after:content-[''] after:absolute after:top-[2px] after:left-[2px] after:bg-white after:border-gray-300 after:border after:rounded-full after:h-5 after:w-5 after:transition-all peer-checked:bg-green-500 shadow-inner"></div>
                </label>
                
                <div class="flex space-x-1 glass p-1 rounded-xl text-xs font-bold">
                    <button onclick="switchTab('scan')" id="tab-scan" class="px-3 py-1.5 rounded-lg transition-all bg-green-500 text-white">Quét Mã</button>
                    <button onclick="switchTab('list')" id="tab-list" class="px-3 py-1.5 rounded-lg transition-all text-gray-500 dark:text-gray-400 hover:text-gray-900 dark:hover:text-white">Danh Sách</button>
                    <button onclick="switchTab('history')" id="tab-history" class="px-3 py-1.5 rounded-lg transition-all text-gray-500 dark:text-gray-400 hover:text-gray-900 dark:hover:text-white">Lịch Sử</button>
                </div>
            </div>
        </div>
    </header>

    <main class="max-w-7xl mx-auto p-4 z-10 relative">

        <!-- TAB 1: QUÉT MÃ QR -->
        <section id="section-scan" class="max-w-md mx-auto space-y-4">
            <div class="glass p-3 rounded-2xl flex items-center space-x-3 shadow-sm">
                <span class="text-xs font-bold sub-text whitespace-nowrap">Người thực hiện:</span>
                <input type="text" id="collaborator-name" placeholder="Nhập tên CTV..." 
                       class="input-glass text-sm font-semibold px-3 py-2 rounded-xl flex-1 focus:outline-none focus:ring-2 focus:ring-green-500"
                       onchange="localStorage.setItem('ctv_name', this.value)">
            </div>

            <!-- Scanner Frame -->
            <div id="scanner-container" class="glass rounded-2xl p-4 shadow-xl border-2 border-green-500/30 text-center">
                <div id="reader" class="w-full"></div>
            </div>

            <!-- Result Card -->
            <div id="result-card" class="hidden glass p-5 rounded-2xl shadow-2xl transition-all duration-300">
                <div class="flex items-center space-x-4 mb-4">
                    <div class="relative">
                        <img id="avatar" src="" class="w-16 h-16 rounded-full object-cover border-2 border-green-400 shadow">
                        <span id="stt-badge" class="absolute -bottom-1 -right-2 bg-green-500 text-white text-[11px] font-extrabold px-2 py-0.5 rounded-md shadow-md"></span>
                    </div>
                    <div class="flex-1">
                        <h2 id="name" class="text-xl font-bold dark-text uppercase"></h2>
                        <p id="address" class="text-xs sub-text mt-1 font-medium"></p>
                        <p id="job" class="text-xs sub-text mt-0.5"></p>
                        <p id="notes" class="text-xs text-yellow-600 dark:text-yellow-400 mt-1 italic font-semibold"></p>
                    </div>
                </div>

                <div id="status-box" class="p-4 rounded-xl mb-5 text-center text-sm font-bold border uppercase tracking-wide"></div>

                <div class="flex flex-col space-y-3">
                    <button id="btn-checkin" onclick="submitCheckin()" class="w-full bg-gradient-to-r from-emerald-500 to-green-600 hover:scale-[1.02] text-white py-3.5 rounded-xl font-bold shadow-lg transition-all hidden uppercase tracking-wide">
                        Xác Nhận Phát Gạo
                    </button>
                    <button onclick="resumeScan()" class="w-full bg-blue-500 hover:bg-blue-600 text-white py-3.5 rounded-xl font-bold shadow-lg transition-all uppercase tracking-wide">
                        Quét Mã Tiếp Theo
                    </button>
                </div>
            </div>
        </section>

        <!-- TAB 2: DANH SÁCH HỘ DÂN -->

        <section id="section-list" class="hidden space-y-4">

            <div class="flex flex-wrap justify-between items-center gap-3">
                <input type="text" id="search-input" placeholder="Tìm kiếm STT, Tên, Địa chỉ, Ghi chú..." oninput="renderBeneficiariesList()"
                    class="input-glass text-sm px-4 py-3 rounded-xl flex-1 max-w-xs focus:outline-none focus:ring-2 focus:ring-green-500 shadow-sm">
                
                <!-- Ô chọn sắp xếp Tab Danh Sách -->
                <div class="flex items-center space-x-2">
                    <label class="text-xs font-bold sub-text whitespace-nowrap">Sắp xếp:</label>
                    <select id="list-sort-select" onchange="renderBeneficiariesList()" class="input-glass text-sm px-3 py-3 rounded-xl font-semibold focus:outline-none focus:ring-2 focus:ring-green-500 shadow-sm">
                        <option value="stt_asc">STT tăng dần</option>
                        <option value="stt_desc">STT giảm dần</option>
                        <option value="name_asc">Họ tên (A - Z)</option>
                        <option value="address_asc">Địa chỉ (A - Z)</option>
                        <option value="status_unreceived">⏳ Chưa nhận gạo</option>
                        <option value="status_received">✅ Đã nhận gạo</option>
                    </select>
                </div>

                <!-- Các nút thao tác -->
                <div class="flex items-center space-x-2">
                    <button onclick="window.location.href='/api/export/beneficiaries'" class="bg-emerald-600 hover:bg-emerald-700 text-white px-4 py-3 rounded-xl font-bold text-sm flex items-center shadow-md transition-colors gap-1">
                        Xuất file Excel
                    </button>
                    <button onclick="openFormModal()" class="bg-green-500 hover:bg-green-600 text-white px-5 py-3 rounded-xl font-bold text-sm flex items-center shadow-md transition-colors">
                        + Thêm Hộ Mới
                    </button>
                </div>
            </div>

            <div class="glass rounded-xl overflow-hidden shadow-lg border border-gray-200 dark:border-gray-700">
                <div class="overflow-x-auto">
                    <table class="w-full text-left text-sm whitespace-nowrap">
                        <thead class="bg-gray-100 dark:bg-gray-800/80 uppercase text-xs font-bold sub-text">
                            <tr>
                                <th class="px-4 py-4 text-center">STT</th>
                                <th class="px-4 py-4">Họ Tên</th>
                                <th class="px-4 py-4">Địa Chỉ</th>
                                <th class="px-4 py-4">Nghề Nghiệp</th>
                                <th class="px-4 py-4">Ghi Chú</th>
                                <th class="px-4 py-4 text-center">Trạng Thái</th>
                                <th class="px-4 py-4 text-center">Thao Tác</th>
                            </tr>
                        </thead>
                        <tbody id="beneficiaries-tbody" class="divide-y divide-gray-200 dark:divide-gray-700/50 dark-text">
                        </tbody>
                    </table>
                </div>
            </div>
        </section>

        <!-- TAB 3: LỊCH SỬ PHÁT GẠO -->
        <section id="section-history" class="hidden space-y-4">
            <!-- Header Tab Lịch sử -->
            <div class="flex flex-wrap justify-between items-center gap-3 glass px-4 py-3 rounded-xl">
                <div class="flex items-center space-x-3">
                    <h3 class="text-sm font-bold dark-text uppercase">Lịch Sử Nhận Gạo Tháng Này</h3>
                    <span id="history-count" class="text-xs text-white font-semibold bg-green-500 px-3 py-1.5 rounded-lg shadow">0 Lượt</span>
                </div>
                <button onclick="window.location.href='/api/export/history'" class="bg-emerald-600 hover:bg-emerald-700 text-white px-4 py-2 rounded-xl font-bold text-sm flex items-center shadow-md transition-colors gap-1">
                    Xuất file Excel
                </button>
            </div>
            
            <!-- Ô Tìm kiếm & Bộ lọc sắp xếp -->
            <div class="flex flex-wrap justify-between items-center gap-3">
                <input type="text" id="history-search-input" placeholder="Tìm kiếm STT, Tên, Địa chỉ, CTV..." oninput="renderHistoryList()"
                    class="input-glass text-sm px-4 py-2.5 rounded-xl flex-1 max-w-sm focus:outline-none focus:ring-2 focus:ring-green-500 shadow-sm">
                
                <div class="flex items-center space-x-2">
                    <label class="text-xs font-bold sub-text whitespace-nowrap">Sắp xếp:</label>
                    <select id="history-sort-select" onchange="renderHistoryList()" class="input-glass text-sm px-3 py-2.5 rounded-xl font-semibold focus:outline-none focus:ring-2 focus:ring-green-500 shadow-sm">
                        <option value="time_desc">Thời gian nhận (Z - A)</option>
                        <option value="time_asc">Thời gian nhận (A - Z)</option>
                        <option value="stt_asc">STT tăng dần</option>
                        <option value="stt_desc">STT giảm dần</option>
                        <option value="name_asc">Tên (A - Z)</option>
                        <option value="address_asc">Địa chỉ (A - Z)</option>
                    </select>
                </div>
            </div>

            <div class="glass rounded-xl overflow-hidden shadow-lg border border-gray-200 dark:border-gray-700">
                <div class="overflow-x-auto">
                    <table class="w-full text-left text-sm whitespace-nowrap">
                        <thead class="bg-gray-100 dark:bg-gray-800/80 uppercase text-xs font-bold sub-text">
                            <tr>
                                <th class="px-4 py-4 text-center">STT</th>
                                <th class="px-4 py-4">Họ Tên</th>
                                <th class="px-4 py-4">Thông Tin</th>
                                <th class="px-4 py-4 text-center">Tháng Đã Nhận</th>
                                <th class="px-4 py-4">Thời Gian Nhận</th>
                                <th class="px-4 py-4 text-center">Thao Tác</th>
                            </tr>
                        </thead>
                        <tbody id="history-container" class="divide-y divide-gray-200 dark:divide-gray-700/50 dark-text">
                        </tbody>
                    </table>
                </div>
            </div>
        </section>

    </main>

    <!-- MODAL: THÊM / SỬA HỘ DÂN -->
    <div id="modal-form" class="fixed inset-0 z-50 bg-black/60 backdrop-blur-sm hidden flex items-center justify-center p-4">
        <div class="glass w-full max-w-md rounded-2xl p-6 space-y-4 shadow-2xl relative">
            <h3 id="modal-title" class="text-xl font-black text-green-500 uppercase">Thêm Hộ Dân</h3>
            <input type="hidden" id="edit-stt">
            
            <div class="space-y-3">
                <div>
                    <label class="text-xs font-bold sub-text">Họ và Tên (*)</label>
                    <input type="text" id="form-name" class="w-full input-glass px-4 py-2.5 rounded-xl mt-1 focus:outline-none focus:ring-2 focus:ring-green-500">
                </div>
                <div>
                    <label class="text-xs font-bold sub-text">Nơi Ở / Địa Chỉ</label>
                    <input type="text" id="form-address" class="w-full input-glass px-4 py-2.5 rounded-xl mt-1 focus:outline-none focus:ring-2 focus:ring-green-500">
                </div>
                <div>
                    <label class="text-xs font-bold sub-text">Nghề Nghiệp</label>
                    <input type="text" id="form-job" class="w-full input-glass px-4 py-2.5 rounded-xl mt-1 focus:outline-none focus:ring-2 focus:ring-green-500">
                </div>
                <div>
                    <label class="text-xs font-bold sub-text">Ghi Chú</label>
                    <input type="text" id="form-notes" class="w-full input-glass px-4 py-2.5 rounded-xl mt-1 focus:outline-none focus:ring-2 focus:ring-green-500">
                </div>
            </div>

            <div class="flex space-x-3 pt-4">
                <button onclick="closeFormModal()" class="w-1/2 bg-gray-500 hover:bg-gray-600 text-white py-3 rounded-xl font-bold text-sm shadow transition-colors">Đóng</button>
                <button onclick="saveBeneficiary()" class="w-1/2 bg-green-500 hover:bg-green-600 text-white py-3 rounded-xl font-bold text-sm shadow transition-colors">Lưu Dữ Liệu</button>
            </div>
        </div>
    </div>

    <!-- Toast Notification -->
    <div id="toast-container" class="fixed bottom-4 left-1/2 -translate-x-1/2 z-50 flex flex-col space-y-2 w-full max-w-sm pointer-events-none"></div>

    
    
    


    <script>
        let currentStt = null;
        let allBeneficiaries = [];
        let allHistory = []; // <--- THÊM BIẾN NÀY để sort ds
        let html5QrcodeScanner = null;
        let isProcessing = false; // Khiên 1: Cờ chống quét trùng

        // Utility: Xử lý ký tự an toàn tránh đơ JS khi tên có dấu ' hoặc "
        function safeStr(str) {
            if (!str) return '';
            return String(str).replace(/'/g, "\\'").replace(/"/g, '&quot;');
        }

        // Utility: Fetch có Hạn Giờ (Timeout 5 giây) tránh treo khi mất mạng
        async function fetchWithTimeout(resource, options = {}) {
            const { timeout = 5000 } = options;
            const controller = new AbortController();
            const id = setTimeout(() => controller.abort(), timeout);
            try {
                const response = await fetch(resource, { ...options, signal: controller.signal });
                clearTimeout(id);
                return response;
            } catch (error) {
                clearTimeout(id);
                throw error;
            }
        }

        // Quản lý Dark/Light Theme
        function toggleTheme() {
            const html = document.documentElement;
            const toggle = document.getElementById('theme-toggle');
            if (html.classList.contains('dark')) {
                html.classList.remove('dark');
                localStorage.setItem('theme', 'light');
                toggle.checked = false;
            } else {
                html.classList.add('dark');
                localStorage.setItem('theme', 'dark');
                toggle.checked = true;
            }
        }

        document.addEventListener("DOMContentLoaded", () => {
            if (localStorage.getItem('theme') === 'light') {
                document.documentElement.classList.remove('dark');
                document.getElementById('theme-toggle').checked = false;
            } else {
                document.getElementById('theme-toggle').checked = true;
            }
            document.getElementById('collaborator-name').value = localStorage.getItem('ctv_name') || '';
        });

        function showToast(msg, type = 'success') {
            const container = document.getElementById('toast-container');
            const toast = document.createElement('div');
            const bg = type === 'error' ? 'bg-red-500' : 'bg-green-500';
            toast.className = `text-white px-5 py-3 rounded-xl shadow-2xl text-sm font-bold text-center ${bg} transition-all duration-300 z-50`;
            toast.innerText = msg;
            container.appendChild(toast);
            setTimeout(() => { toast.style.opacity = '0'; setTimeout(() => toast.remove(), 300); }, 2500);
        }


        function startScanner() {
            html5QrcodeScanner = new Html5QrcodeScanner("reader", { fps: 10, qrbox: {width: 250, height: 250} }, false);
            html5QrcodeScanner.render(onScanSuccess);
        }

        
        // Đổi Tab & Quản lý trạng thái Camera an toàn
        function switchTab(tab) {
            ['scan', 'list', 'history'].forEach(t => {
                document.getElementById(`section-${t}`).classList.add('hidden');
                document.getElementById(`tab-${t}`).className = 'px-3 py-1.5 rounded-lg transition-all text-gray-500 dark:text-gray-400 hover:text-gray-900 dark:hover:text-white';
            });
            document.getElementById(`section-${tab}`).classList.remove('hidden');
            document.getElementById(`tab-${tab}`).className = 'px-3 py-1.5 rounded-lg transition-all bg-green-500 text-white shadow';

            if (tab === 'list') loadBeneficiaries();
            if (tab === 'history') loadHistory();
            
            try {
                if (tab === 'scan') {
                    // Chỉ bật lại camera nếu đang không có kết quả nào hiển thị trên thẻ result-card
                    const isResultHidden = document.getElementById('result-card').classList.contains('hidden');
                    if (html5QrcodeScanner && html5QrcodeScanner.getState() === 3 && isResultHidden) { 
                        html5QrcodeScanner.resume();
                    } else if (!html5QrcodeScanner) {
                        startScanner();
                    }
                } else {
                    if (html5QrcodeScanner && html5QrcodeScanner.getState() === 2) { 
                        html5QrcodeScanner.pause(true);
                    }
                }
            } catch (err) {
                console.warn("Camera lifecycle warning:", err);
            }
        }

        function resumeScan() {
            isProcessing = false; // Mở khóa cờ
            
            // Lỗi Timer setTimeout: Chỉ thực sự mở lại UI camera nếu người dùng đang ở Tab Scan
            const isScanTabActive = !document.getElementById('section-scan').classList.contains('hidden');
            
            if (isScanTabActive) {
                document.getElementById('result-card').classList.add('hidden');
                document.getElementById('scanner-container').classList.remove('hidden');
                try {
                    if (html5QrcodeScanner && html5QrcodeScanner.getState() === 3) {
                        html5QrcodeScanner.resume();
                    }
                } catch (err) {
                    console.warn("Cannot resume camera:", err);
                }
            } else {
                // Nếu đang ở tab khác thì chỉ ẩn result card đi, không resume camera
                document.getElementById('result-card').classList.add('hidden');
                document.getElementById('scanner-container').classList.remove('hidden');
            }
        }




        async function onScanSuccess(decodedText) {
            if (isProcessing) return;
            isProcessing = true;

            if (html5QrcodeScanner && html5QrcodeScanner.getState() === 2) {
                html5QrcodeScanner.pause(true);
            }
            
            let rawCode = decodedText.trim();
            if (!rawCode) {
                showToast("Mã QR rỗng!", "error");
                setTimeout(resumeScan, 1500);
                return;
            }
            
            document.getElementById('scanner-container').classList.add('hidden');
            
            try {
                // Gửi nguyên chuỗi mã quét được lên server xử lý
                const res = await fetchWithTimeout(`/api/beneficiary/${encodeURIComponent(rawCode)}`, { timeout: 4000 });
                const data = await res.json();

                if (data.detail) { 
                    showToast(data.detail, 'error'); 
                    setTimeout(resumeScan, 2000);
                    return; 
                }
                
                currentStt = data.stt;
                document.getElementById('name').innerText = data.full_name;
                document.getElementById('stt-badge').innerText = `STT #${data.stt}`;
                document.getElementById('address').innerText = `📍 ${data.address || 'Đang cập nhật'}`;
                document.getElementById('job').innerText = `💼 Nghề nghiệp: ${data.job || '---'}`;
                document.getElementById('notes').innerText = data.notes ? `📝 Ghi chú: ${data.notes}` : '';
                document.getElementById('avatar').src = data.avatar_url;
                
                const statusBox = document.getElementById('status-box');
                const btn = document.getElementById('btn-checkin');
                
                if (data.already_received) {
                    statusBox.className = "p-4 rounded-xl mb-5 text-center text-sm font-bold border uppercase tracking-wide bg-red-50 dark:bg-red-900/30 text-red-600 dark:text-red-400 border-red-500/50 shadow-inner";
                    statusBox.innerHTML = `ĐÃ NHẬN THÁNG NÀY<br><span class="text-xs opacity-90 capitalize font-medium mt-1 inline-block">Lúc: ${data.received_at} (${data.collaborator_name})</span>`;
                    btn.classList.add('hidden');
                } else {
                    statusBox.className = "p-4 rounded-xl mb-5 text-center text-sm font-bold border uppercase tracking-wide bg-green-50 dark:bg-green-900/30 text-green-600 dark:text-green-400 border-green-500/50 shadow-inner";
                    statusBox.innerText = "CHƯA NHẬN GẠO";
                    btn.classList.remove('hidden');
                    btn.disabled = false;
                }
                
                document.getElementById('result-card').classList.remove('hidden');
            } catch (err) {
                showToast("Lỗi kết nối hoặc hết thời gian phản hồi!", "error");
                setTimeout(resumeScan, 1500);
            }
        }




        

        



        // Xác nhận phát gạo (Có chống spam click)
        async function submitCheckin() {
            const btn = document.getElementById('btn-checkin');
            btn.disabled = true; // Khóa nút ngay lập tức
            
            const ctvInput = document.getElementById('collaborator-name').value.trim();
            const ctv = ctvInput || 'CTV';
            if (ctvInput) localStorage.setItem('ctv_name', ctvInput);

            try {
                const res = await fetchWithTimeout('/api/checkin', {
                    method: 'POST',
                    headers: {'Content-Type': 'application/json'},
                    body: JSON.stringify({stt: currentStt, collaborator_name: ctv}),
                    timeout: 5000
                });
                const data = await res.json();
                
                if (data.detail) showToast(data.detail, 'error');
                else {
                    showToast(data.message);
                    resumeScan();
                }
            } catch (e) {
                showToast("Lỗi mạng khi xác nhận!", "error");
                btn.disabled = false;
            }
        }

        async function manualCheckin(stt) {
            const ctvInput = document.getElementById('collaborator-name').value.trim();
            const ctv = ctvInput || 'CTV';
            if (ctvInput) localStorage.setItem('ctv_name', ctvInput);

            try {
                const res = await fetchWithTimeout('/api/checkin', {
                    method: 'POST',
                    headers: {'Content-Type': 'application/json'},
                    body: JSON.stringify({stt: stt, collaborator_name: ctv})
                });
                const data = await res.json();
                if (data.detail) showToast(data.detail, 'error');
                else {
                    showToast(data.message);
                    loadBeneficiaries(); 
                }
            } catch (e) {
                showToast("Lỗi kết nối!", "error");
            }
        }


        async function undoCheckin(stt) {
            // Tự động tìm tên dựa vào stt từ mảng đang có sẵn
            let person = allBeneficiaries.find(b => b.stt === stt) || allHistory.find(h => h.stt === stt);
            let name = person ? person.full_name : "người này";

            if (!confirm(`⚠️ Bạn có chắc chắn muốn HỦY xác nhận phát gạo của hộ: ${name} (STT: #${stt}) không?`)) return;
            try {
                const res = await fetchWithTimeout(`/api/checkin/${stt}`, { method: 'DELETE' });
                const data = await res.json();
                if (data.detail) showToast(data.detail, 'error');
                else {
                    showToast(`Đã thu hồi lịch sử của STT #${stt} thành công!`);
                    if (!document.getElementById('section-list').classList.contains('hidden')) loadBeneficiaries();
                    if (!document.getElementById('section-history').classList.contains('hidden')) loadHistory();
                }
            } catch (e) {
                showToast("Lỗi khi gửi yêu cầu thu hồi!", "error");
            }
        }



        // Render danh sách người dân (Sử dụng createElement an toàn của V2)
        function loadBeneficiaries() {
            fetch('/api/beneficiaries')
            .then(res => res.json())
            .then(data => {
                allBeneficiaries = data;
                renderBeneficiariesList();
            }).catch(() => showToast("Không tải được danh sách!", "error"));
        }

        // Thay thế hàm renderBeneficiariesList() cũ:
        function renderBeneficiariesList() {
            const search = document.getElementById('search-input').value.toLowerCase().trim();
            const sortMode = document.getElementById('list-sort-select')?.value || 'stt_asc';
            const tbody = document.getElementById('beneficiaries-tbody');
            tbody.innerHTML = '';
            
            // 1. Lọc theo từ khóa
            let filtered = allBeneficiaries.filter(b => 
                String(b.stt).includes(search) || 
                (b.full_name && b.full_name.toLowerCase().includes(search)) || 
                (b.address && b.address.toLowerCase().includes(search)) ||
                (b.notes && b.notes.toLowerCase().includes(search))
            );

            // 2. Sắp xếp danh sách
            filtered.sort((a, b) => {
                if (sortMode === 'stt_asc') return a.stt - b.stt;
                if (sortMode === 'stt_desc') return b.stt - a.stt;
                if (sortMode === 'name_asc') return a.full_name.localeCompare(b.full_name, 'vi');
                if (sortMode === 'address_asc') return (a.address || '').localeCompare(b.address || '', 'vi');
                if (sortMode === 'status_unreceived') return (a.received_at ? 1 : 0) - (b.received_at ? 1 : 0);
                if (sortMode === 'status_received') return (b.received_at ? 1 : 0) - (a.received_at ? 1 : 0);
                return 0;
            });

            // 3. Render giao diện
            filtered.forEach(b => {
                const tr = document.createElement('tr');
                tr.className = "hover:bg-gray-50 dark:hover:bg-gray-800/50 transition-colors";
                
                const isReceived = !!b.received_at;
                const statusBadge = isReceived 
                    ? `<span class="bg-green-100 dark:bg-green-900/50 text-green-700 dark:text-green-400 px-2 py-1 rounded text-xs font-bold border border-green-200 dark:border-green-800">Đã nhận</span>`
                    : `<span class="bg-yellow-100 dark:bg-yellow-900/50 text-yellow-700 dark:text-yellow-400 px-2 py-1 rounded text-xs font-bold border border-yellow-200 dark:border-yellow-800">Chưa nhận</span>`;
                
                const actionBtn = isReceived
                    ? `<button onclick="undoCheckin(${b.stt})" class="text-red-500 hover:text-red-700 font-bold text-xs p-1 transition">Hủy Phát</button>`
                    : `<button onclick="manualCheckin(${b.stt})" class="text-green-500 hover:text-green-700 font-bold text-xs p-1 transition">Phát Gạo</button>`;

                tr.innerHTML = `
                    <td class="px-4 py-3 text-center font-black">#${b.stt}</td>
                    <td class="px-4 py-3 font-bold uppercase">${b.full_name}</td>
                    <td class="px-4 py-3 text-xs">${b.address || '---'}</td>
                    <td class="px-4 py-3 text-xs">${b.job || '---'}</td>
                    <td class="px-4 py-3 text-xs italic text-yellow-600 dark:text-yellow-400 font-medium">${b.notes || ''}</td>
                    <td class="px-4 py-3 text-center">${statusBadge}</td>
                    <td class="px-4 py-3 text-center space-x-2">
                        ${actionBtn}
                        <button onclick="openFormModal(${b.stt})" class="text-blue-500 hover:text-blue-700 font-bold text-xs p-1 transition">Sửa</button>
                    </td>
                `;
                tbody.appendChild(tr);
            });
        }




        // Render Lịch Sử Phát Gạo (Kết hợp render Badge Tháng đẹp của V2)
        function loadHistory() {
            fetch('/api/checkins/history')
            .then(res => res.json())
            .then(data => {
                allHistory = data;
                renderHistoryList();
            }).catch(() => showToast("Lỗi khi tải lịch sử!", "error"));
        }

        function renderHistoryList() {
            const search = (document.getElementById('history-search-input')?.value || '').toLowerCase().trim();
            const sortMode = document.getElementById('history-sort-select')?.value || 'time_desc';
            
            // 1. Lọc theo từ khóa tìm kiếm
            let filtered = allHistory.filter(h => 
                String(h.stt).includes(search) || 
                (h.full_name && h.full_name.toLowerCase().includes(search)) || 
                (h.address && h.address.toLowerCase().includes(search)) ||
                (h.notes && h.notes.toLowerCase().includes(search)) ||
                (h.collaborator_name && h.collaborator_name.toLowerCase().includes(search))
            );

            // 2. Sắp xếp danh sách lịch sử
            filtered.sort((a, b) => {
                if (sortMode === 'time_desc') return b.id - a.id;
                if (sortMode === 'time_asc') return a.id - b.id;
                if (sortMode === 'stt_asc') return a.stt - b.stt;
                if (sortMode === 'stt_desc') return b.stt - a.stt;
                if (sortMode === 'name_asc') return a.full_name.localeCompare(b.full_name, 'vi');
                if (sortMode === 'address_asc') return (a.address || '').localeCompare(b.address || '', 'vi');
                return 0;
            });

            // 3. Cập nhật đếm số lượng
            document.getElementById('history-count').innerText = `${filtered.length} / ${allHistory.length} Lượt`;
            const tbody = document.getElementById('history-container');
            tbody.innerHTML = '';
            
            // 4. Render các dòng của bảng
            filtered.forEach(h => {
                const tr = document.createElement('tr');
                tr.className = "hover:bg-gray-50 dark:hover:bg-gray-800/50 transition-colors";

                let monthsHtml = '-';
                if (h.months_received) {
                    monthsHtml = h.months_received.split(',')
                        .map(m => `<span class="bg-blue-100 text-blue-700 dark:bg-blue-900/40 dark:text-blue-300 px-1.5 py-0.5 rounded text-[10px] font-bold mx-0.5 inline-block">Tháng ${m.trim()}</span>`)
                        .join('');
                }

                tr.innerHTML = `
                    <td class="px-4 py-3 text-center font-black">#${h.stt}</td>
                    <td class="px-4 py-3 font-bold uppercase">${h.full_name}</td>
                    <td class="px-4 py-3 text-xs">
                        📍 ${h.address || '---'}<br>
                        <span class="italic text-yellow-600 dark:text-yellow-400 font-medium">📝 ${h.notes || '---'}</span>
                    </td>
                    <td class="px-4 py-3 text-center">${monthsHtml}</td>
                    <td class="px-4 py-3 text-xs font-semibold">
                        🕒 ${h.created_at}<br>
                        <span class="opacity-70 font-normal">👤 CTV: ${h.collaborator_name}</span>
                    </td>
                    <td class="px-4 py-3 text-center">
                        <button onclick="undoCheckin(${h.stt})" class="text-red-500 hover:text-red-700 font-bold text-xs bg-red-100 dark:bg-red-900/30 px-3 py-1.5 rounded-lg transition">Thu hồi</button>
                    </td>
                `;
                tbody.appendChild(tr);
            });
        }


        function openFormModal(stt = null) {
            const modal = document.getElementById('modal-form');
            const title = document.getElementById('modal-title');
            
            if (stt) {
                title.innerText = "Cập Nhật Thông Tin";
                document.getElementById('edit-stt').value = stt;
                const b = allBeneficiaries.find(x => x.stt === stt);
                if (b) {
                    document.getElementById('form-name').value = b.full_name;
                    document.getElementById('form-address').value = b.address || '';
                    document.getElementById('form-job').value = b.job || '';
                    document.getElementById('form-notes').value = b.notes || '';
                }
            } else {
                title.innerText = "Thêm Hộ Dân Mới";
                document.getElementById('edit-stt').value = "";
                document.getElementById('form-name').value = "";
                document.getElementById('form-address').value = "";
                document.getElementById('form-job').value = "";
                document.getElementById('form-notes').value = "";
            }
            modal.classList.remove('hidden');
        }

        function closeFormModal() {
            document.getElementById('modal-form').classList.add('hidden');
        }

        async function saveBeneficiary() {
            const stt = document.getElementById('edit-stt').value;
            const payload = {
                full_name: document.getElementById('form-name').value.trim(),
                address: document.getElementById('form-address').value.trim(),
                job: document.getElementById('form-job').value.trim(),
                notes: document.getElementById('form-notes').value.trim()
            };

            if (!payload.full_name) {
                showToast("Vui lòng nhập Họ và Tên!", "error");
                return;
            }

            const method = stt ? 'PUT' : 'POST';
            const url = stt ? `/api/beneficiary/${stt}` : '/api/beneficiary';

            try {
                const res = await fetchWithTimeout(url, {
                    method: method,
                    headers: {'Content-Type': 'application/json'},
                    body: JSON.stringify(payload)
                });
                const resData = await res.json();
                if (resData.detail) showToast(resData.detail, "error");
                else {
                    showToast(resData.message);
                    closeFormModal();
                    if (!document.getElementById('section-list').classList.contains('hidden')) {
                        loadBeneficiaries();
                    }
                }
            } catch (e) {
                showToast("Lỗi kết nối khi lưu!", "error");
            }
        }

        // Mặc định bật tab Scan khi ứng dụng sẵn sàng
        window.onload = () => switchTab('scan');
    </script>

</body>
</html>
"""

@app.get("/")
def get_home():
    return HTMLResponse(content=HTML_TEMPLATE)




# Chạy ứng dụng: uvicorn app-phat-gao:app --reload --host 0.0.0.0 --port 8000
# 1. uvicorn app-phat-gao:app --reload
# 2. .\cloudflared.exe tunnel --url http://localhost:8000